# pptx_generator.py - PowerPoint generator for the pitch deck
"""
Generates a 16:9 widescreen PowerPoint (.pptx) file from the presentation data
structure defined in backend/default_data.py.
The design follows a cybersecurity dark theme with navy background and cyan accents.
"""

import io
from typing import Dict, List

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE

# Theme colors
NAVY = RGBColor(10, 20, 60)          # Background color
CYAN = RGBColor(0, 200, 255)         # Accent color for titles
WHITE = RGBColor(255, 255, 255)


def _add_title_slide(prs: Presentation, title: str, subtitle: str) -> None:
    """Add the title slide with the main title and subtitle."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    # Set background
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY

    left = Inches(0.8)
    top = Inches(1.5)
    width = Inches(12)
    height = Inches(1.5)
    # Title
    title_box = slide.shapes.add_textbox(left, top, width, height)
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(48)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = CYAN
    # Subtitle
    sub_box = slide.shapes.add_textbox(left, top + Inches(2), width, Inches(1))
    sub_tf = sub_box.text_frame
    sub_tf.text = subtitle
    sub_tf.paragraphs[0].font.size = Pt(28)
    sub_tf.paragraphs[0].font.color.rgb = WHITE


def _add_content_slide(prs: Presentation, title: str, content: List[str]) -> None:
    """Add a generic content slide with a title and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY

    left = Inches(0.8)
    top = Inches(0.8)
    width = Inches(12)
    # Title
    title_box = slide.shapes.add_textbox(left, top, width, Inches(1))
    tf = title_box.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(36)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = CYAN

    # Content bullets
    body_top = top + Inches(1.5)
    body_box = slide.shapes.add_textbox(left, body_top, width, Inches(4))
    body_tf = body_box.text_frame
    for point in content:
        p = body_tf.add_paragraph()
        p.text = point
        p.font.size = Pt(24)
        p.font.color.rgb = WHITE
        p.level = 0


def generate_pitch_deck_pptx(presentation_data: Dict) -> bytes:
    """Generate a PPTX file from the provided presentation dictionary.
    Returns the binary PPTX data.
    """
    prs = Presentation()
    # Set slide size to 16:9 (13.33 x 7.5 inches)
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    metadata = presentation_data.get("metadata", {})
    # Title slide using metadata
    _add_title_slide(
        prs,
        metadata.get("project_name", "Presentation"),
        metadata.get("subtitle", ""),
    )

    # Slides
    for slide in presentation_data.get("slides", []):
        _add_content_slide(prs, slide.get("title", ""), slide.get("content", []))

    # Save to in‑memory buffer
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()

# The module can also be executed directly for a quick test
if __name__ == "__main__":
    from backend.default_data import get_default_presentation
    data = get_default_presentation()
    pptx_bytes = generate_pitch_deck_pptx(data)
    with open("test_deck.pptx", "wb") as f:
        f.write(pptx_bytes)
    print("Test PPTX generated: test_deck.pptx")
